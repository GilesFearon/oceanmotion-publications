#!/usr/bin/env python3
"""
Ocean Motion Analytics — ports and harbours pitch deck, focused on under-keel
clearance.

Built on top of the desal deck as it was actually sent, not as make_desal_deck.py
generates it: oma-pitch-desal-2026-09-23.pptx carries hand edits made in
PowerPoint (reworded headlines, re-laid chain slides, the dashboard and
with-your-data slides) that the generator does not know about. Regenerating
from code would silently throw those away. So this script opens that file,
keeps the slides that carry over, rewords them run by run so their formatting
survives, drops the turbidity slides, and adds the water level slides with the
same deckkit helpers everything else was built with.

The consequence: if the desal deck is edited by hand again, this deck will pick
the edits up on its next build -- but shape ids below (KEEP / the retext calls)
are that file's, so a structural edit there can break a lookup here. Every
lookup fails loudly rather than skipping.

Story, following the desal deck's thread:
  01  how the model works        (kept: inputs, grid, circulation, waves)
  02  tested against tide gauges UKC as a water level budget; the predicted
                                 tide at Salmiya and Majis; the non-tidal
                                 residual at both
  03  water levels at your port  the modelled residual at Khalifa Port -- no
                                 gauge, said so on the slide -- and a week of
                                 spring lows where the total water level went
                                 under the tide table
  04  what we'd build for you    design hindcast / daily UKC forecast /
                                 with your data / climate outlook

Figures come from make_water_levels.py (somisana_croco env). Numbers quoted on
the slides are the ones that script prints; if the data changes, rerun it and
check them here.

Runs in the base environment (needs python-pptx):
    python3 make_ports_deck.py
Output: oma-pitch-ports.pptx
"""

import argparse
import copy
import os
from datetime import date

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.text.text import _Paragraph

import deckkit as dk
from deckkit import (ABYSS, ACCENT, CYAN, CYAN_LT, INK, INK_2, INK_3, NAVY,
                     NAVY_MID, PAPER, PAPER_2, RULE,
                     F_BODY, F_DISPLAY, F_MONO, EMU_W, EMU_H, MARGIN)
import make_chain_deck as chain
import make_desal_deck as desal

BASE = os.path.join(dk.HERE, "oma-pitch-desal-2026-09-23.pptx")
ASSET = dk.ASSET_DIR


def asset(name):
    path = os.path.join(ASSET, name)
    if not os.path.exists(path):
        raise SystemExit(f"{path} missing — run make_water_levels.py --all "
                         "in the somisana_croco env first")
    return path


# ------------------------------------------------------------------ retext
# The base deck has been through Google Slides and PowerPoint, so its runs are
# fragmented ('W' + 'ater quality' + ...) and formatting sits on whichever run
# survived. These helpers replace text while borrowing a named run's <a:rPr>,
# so the new words wear the old formatting.

def shape(s, sid):
    for sh in s.shapes:
        if sh.shape_id == sid:
            return sh
    raise SystemExit(f"shape {sid} not found on slide — has {BASE} been "
                     "restructured? Update the ids in make_ports_deck.py.")


def set_runs(p, segs):
    """Rewrite one paragraph. segs is [(text, template_run_index)]; a '\\n' in
    the text becomes a line break carrying the same formatting."""
    tmpl = [copy.deepcopy(r) for r in p._p.findall(qn("a:r"))]
    for el in list(p._p):
        if el.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p._p.remove(el)
    end = p._p.find(qn("a:endParaRPr"))

    def put(el):
        if end is not None:
            end.addprevious(el)
        else:
            p._p.append(el)

    for text, ti in segs:
        src = tmpl[ti]
        for k, part in enumerate(text.split("\n")):
            if k:
                br = p._p.makeelement(qn("a:br"), {})
                rpr = src.find(qn("a:rPr"))
                if rpr is not None:
                    br.append(copy.deepcopy(rpr))
                put(br)
            if part:
                r = copy.deepcopy(src)
                r.find(qn("a:t")).text = part
                put(r)


def set_paras(sh, items):
    """Rewrite a text frame paragraph by paragraph. Paragraph i borrows the
    i-th existing paragraph's formatting, or the last one's past the end."""
    body = sh.text_frame._txBody
    olds = body.findall(qn("a:p"))
    tmpl = [copy.deepcopy(p) for p in olds]
    for p in olds:
        body.remove(p)
    nruns = [len(p.findall(qn("a:r"))) for p in tmpl]
    for i, segs in enumerate(items):
        # the i-th paragraph if it has the runs this one borrows from, else
        # the nearest that does -- a value-only spec line has one run, and a
        # key/value row must not take its key's formatting from it
        need = max(ti for _, ti in segs) + 1
        j = min(i, len(tmpl) - 1)
        if nruns[j] < need:
            ok = [n for n in range(len(tmpl)) if nruns[n] >= need]
            if not ok:
                raise SystemExit(f"no paragraph in {sh.name} has {need} runs")
            j = min(ok, key=lambda n: abs(n - j))
        el = copy.deepcopy(tmpl[j])
        body.append(el)
        set_runs(_Paragraph(el, sh.text_frame), segs)


def dash_item(text):
    """A '—  text' bullet as the base deck writes it: dash run, then text run."""
    return [("—  ", 0), (text, 1)]


def spec_item(key, val):
    return [(f"{key.upper()}  ", 0), (val, 1)] if key else [(val, -1)]


def retext(s, sid, text, run=0):
    """Single-paragraph shorthand."""
    set_paras(shape(s, sid), [[(text, run)]])


# ------------------------------------------------------------------ structure
def drop_slide(prs, s):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is s.part:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
            return
    raise SystemExit("slide not in deck")


def reorder(prs, order):
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(sid.rId): sid for sid in lst}
    ids = [by_part[s.part] for s in order]
    assert len(ids) == len(lst), f"order has {len(ids)} of {len(lst)} slides"
    for sid in list(lst):
        lst.remove(sid)
    for sid in ids:
        lst.append(sid)


def new_slide(prs, ground=PAPER):
    # layout 0 is BLANK in the base deck; its date/footer/number placeholders
    # are not cloned onto new slides, so this comes out genuinely empty
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    dk.bg(s, ground)
    return s


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ------------------------------------------------------------------ pieces
# The progress rail as the base deck's chain slides now draw it -- five ticks,
# labels at 0.55 in, not the six of make_chain_deck.RAIL -- with the turbidity
# stop renamed. Measured off slide 5 of the base deck.
RAIL = ["inputs", "grid", "circulation", "waves", "water level"]


def rail(s, active):
    for i, label in enumerate(RAIL):
        x = Inches(1.02) + i * Inches(1.14)
        live = i == active
        dk.rect(s, x, Inches(0.42), Inches(0.92), Pt(2.2),
                ACCENT if live else RULE)
        tf = dk.box(s, x, Inches(0.55), Inches(1.14), Inches(0.13))
        dk.run(tf.paragraphs[0], label, font=F_MONO, size=8,
               color=INK if live else INK_3, track=70, caps=True)


def divider(prs, tag, text):
    s = new_slide(prs, NAVY)
    dk.eyebrow(s, MARGIN, Inches(3.0), tag, color=CYAN)
    desal.headline(s, MARGIN, Inches(3.5), Inches(11.4), text, None, size=58,
                   color=PAPER)
    return s


def stat(s, left, top, big, label, width=Inches(2.5), size=54):
    """One big number and its label -- the metrics row, one at a time, so a
    few can sit beside a headline instead of taking a band of their own."""
    tf = dk.box(s, left, top, width, Inches(0.95))
    dk.run(tf.paragraphs[0], big, font=F_DISPLAY, size=size, color=ACCENT)
    tf = dk.box(s, left, top + Inches(size / 72 * 1.05), width, Inches(0.6))
    p = tf.paragraphs[0]
    p.line_spacing = 1.3
    dk.run(p, label, font=F_MONO, size=9.5, color=INK_2, track=100, caps=True)


def line(s, x0, y0, x1, y1, color, width=Pt(1.25), dash=None, heads=()):
    ln_ = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *(dk._emu(v) for v in
                                                          (x0, y0, x1, y1)))
    ln_.line.color.rgb = color
    ln_.line.width = width
    if dash:
        ln_.line.dash_style = dash
    ln = ln_.line._get_or_add_ln()
    for end in heads:
        ln.append(ln.makeelement(qn(f"a:{end}End"),
                                 {"type": "triangle", "w": "sm", "len": "sm"}))
    return ln_


def label(s, left, top, width, text, color=INK_2, size=9, align=PP_ALIGN.LEFT,
          font=F_MONO, caps=False):
    tf = dk.box(s, left, top, width, Inches(0.26), anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = align
    dk.run(p, text, font=font, size=size, color=color, track=60, caps=caps)
    return tf


# ------------------------------------------------------------------ 02
def ukc_slide(prs):
    """UKC as a water level budget, drawn as native shapes so every level and
    label can be moved in the room.

    Drawn as a set-down: the water sits below the tide-table level, so the
    clearance the passage plan counted on is not all there. That is the one
    case the rest of the deck is about."""
    s = new_slide(prs)
    rail(s, 4)
    desal.headline(s, MARGIN, Inches(1.20), Inches(11.4),
                   "Under-keel clearance is a water level budget.", size=34)

    lx, lw = MARGIN, Inches(1.45)                  # level labels, right-aligned
    x0, x1 = Inches(2.50), Inches(6.20)            # water
    y_tt, y_ws, y_cd = Inches(3.00), Inches(3.42), Inches(4.20)
    y_keel, y_sb, y_bed = Inches(5.40), Inches(6.02), Inches(6.42)

    dk.rect(s, x0, y_ws, x1 - x0, y_sb - y_ws, RGBColor(0xd6, 0xea, 0xee))
    dk.rect(s, x0, y_sb, x1 - x0, y_bed - y_sb, RGBColor(0xe0, 0xd2, 0xae))

    line(s, x0, y_tt, x1, y_tt, INK_3, Pt(1.25), MSO_LINE_DASH_STYLE.DASH)
    line(s, x0, y_ws, x1, y_ws, NAVY, Pt(1.75))
    line(s, x0, y_cd, x1, y_cd, INK_3, Pt(0.9), MSO_LINE_DASH_STYLE.ROUND_DOT)

    for y, text, c in ((y_tt, "tide table", INK_3),
                       (y_ws, "actual water", NAVY),
                       (y_cd, "chart datum", INK_3),
                       (y_sb, "dredged depth", INK_2)):
        label(s, lx, y - Inches(0.13), lw, text, color=c, align=PP_ALIGN.RIGHT)

    # the hull: an inverted trapezoid, drawn over the levels it sits in
    hull = s.shapes.add_shape(MSO_SHAPE.FLOWCHART_MANUAL_OPERATION,
                              Inches(3.00), Inches(2.62), Inches(2.10),
                              y_keel - Inches(2.62))
    hull.fill.solid(); hull.fill.fore_color.rgb = NAVY_MID
    hull.line.fill.background()
    dk.flatten(hull)

    # clearance, under the keel
    cx = Inches(4.05)
    line(s, cx, y_keel + Pt(2), cx, y_sb - Pt(2), ACCENT, Pt(1.5),
         heads=("head", "tail"))
    label(s, cx + Inches(0.12), y_keel + Inches(0.18), Inches(1.9),
          "under-keel clearance", color=ACCENT, size=9)

    # the residual, beside the ship
    rx = Inches(5.62)
    line(s, rx, y_tt + Pt(2), rx, y_ws - Pt(2), ACCENT, Pt(1.5),
         heads=("head", "tail"))
    label(s, rx + Inches(0.10), y_tt - Inches(0.44), Inches(1.3),
          "residual", color=ACCENT, size=9)

    desal.bullets(s, Inches(7.10), Inches(2.72), Inches(5.3), [
        "Tide tables give you the astronomical tide, years ahead.",
        "They say nothing about the non-tidal residual: wind set-up and "
        "set-down, air pressure, and the slow sloshing of the Gulf itself.",
        "A set-down at low water is clearance the passage plan counted on and "
        "does not have.",
        "The model forecasts tide and residual together, hourly, with the "
        "waves that drive vessel motion alongside.",
    ], size=14)
    notes(s, "The drawing is a set-down: the water is below the tide-table "
             "level, so the UKC in the passage plan is not all there. Squat and "
             "wave response are the other two terms of the budget; the wave "
             "model on the previous slide covers the second.")
    return s


def tide_slide(prs):
    """The predicted tide as context, not as the pitch: a port's local
    forecast takes its tide from the gauge's own harmonics, so what this slide
    has to show is that the model's tide is sound, over a whole month, with
    the gauge records' lengths stated."""
    s = new_slide(prs)
    rail(s, 4)
    desal.headline(s, MARGIN, Inches(0.92), Inches(6.2),
                   "The predicted tide,\nchecked at two gauges.", size=30)
    chain.specs(s, Inches(7.35), Inches(1.00), Inches(5.3),
                [("Majis", "6 Feb 2023 – 24 Mar 2024  ·  14 months"),
                 ("Salmiya", "15 Jun 2023 – 26 Mar 2024  ·  9 months"),
                 ("sampling", "hourly"),
                 ("timing", "model ~25 min late at Majis,"),
                 ("", "1 – 1.5 h late at Salmiya")],
                "gauge records")
    for k, name in enumerate(("Majis", "Salmiya")):
        top = Inches(2.05) + k * Inches(2.35)
        label(s, MARGIN, top, Inches(8), f"{GAUGE[name]}  ·  August 2023",
              color=INK_2, size=9.5)
        dk.picture(s, asset(f"wl-tide-{name.lower()}.png"), MARGIN,
                   top + Inches(0.20), width=Inches(11.0))
    desal.caption(s, MARGIN, Inches(6.88), Inches(11.5),
                  "operationally the local forecast takes its tide from the "
                  "gauge's own harmonics — the residual is what the model brings")
    notes(s, "Model and gauge harmonics put through the same utide fit over "
             "each gauge's whole record; r and RMSE on the figures are over "
             "that whole record, August 2023 is only what is drawn. The model "
             "constituents are the operational ones, fitted over each gauge's "
             "period exactly as the gauge is. Majis: "
             "amplitudes within 2.1 cm on M2 S2 N2 K1 O1, and the model lags "
             "by the same ~25 min on every constituent - looks like a "
             "timestamp convention on the hourly averages, to be checked. "
             "Salmiya: M2 -5.5 cm, O1 +7 cm, lag 61 min (M2) to 95 min (K1); "
             "the RMSE of 38 cm is mostly that timing, accumulated along the "
             "tide's path from Hormuz to the head of the Gulf.")
    return s


GAUGE = {"Salmiya": "Salmiya, Kuwait", "Majis": "Majis, Oman"}


def residual_slide(prs, name, head, stats, cap, note_text):
    s = new_slide(prs)
    rail(s, 4)
    desal.headline(s, MARGIN, Inches(1.00), Inches(5.6), head, size=30)
    for k, (big, lbl) in enumerate(stats):
        stat(s, Inches(6.75) + k * Inches(2.0), Inches(0.95), big, lbl,
             width=Inches(1.9), size=40)
    dk.picture(s, asset(f"wl-residual-{name.lower()}.png"), MARGIN,
               Inches(2.72), width=Inches(11.0))
    tf = dk.box(s, MARGIN, Inches(6.42), Inches(11.5), Inches(0.9))
    p = tf.paragraphs[0]
    p.line_spacing = 1.3
    dk.run(p, cap, font=F_BODY, size=12.5, color=INK_2, italic=True)
    notes(s, note_text)
    return s


RESID_METHOD = ("Residual only. Gauge, CROCO and MERCATOR each put through the "
                "same harmonic analysis over the gauge record period (utide, "
                "trend=False - the record is too short to fit a trend), exactly "
                "as the operational forecast does. MERCATOR has no tide, but "
                "the fit removes the seasonal constituents the gauge tide "
                "already carries, so they are not counted twice. MERCATOR is "
                "not forced by air pressure, so the inverse barometer is added "
                "to it first (ERA5 here, the forecast's own forcing pressure "
                "operationally). These are the "
                "numbers in croco_residual_stats.nc and "
                "mercator_residual_stats.nc, which set the forecast's "
                "confidence bands. Operationally the ensemble grows: CROCO "
                "forced by GFS and by ECMWF, alongside MERCATOR.")


def salmiya_slide(prs):
    return residual_slide(
        prs, "Salmiya", "Salmiya, Kuwait: the surge,\nevent by event.",
        [("0.88", "CROCO\ncorrelation\nRMSE 8.7 cm"),
         ("0.93", "MERCATOR\ncorrelation\nRMSE 6.9 cm"),
         ("0.94", "mean of the two\ncorrelation\nRMSE 6.5 cm")],
        "Residual only: gauge, our model and MERCATOR (with the air-pressure "
        "response added) put through one harmonic analysis over the gauge "
        "record, as the live forecast does; 15 Jun 2023 – 26 Mar 2024. Some "
        "events are caught best by our "
        "model, others by MERCATOR, and the mean of the two beats either — "
        "operationally the ensemble grows further, with our model forced by "
        "both GFS and ECMWF.",
        "Salmiya: 6 840 hourly pairs. CROCO r 0.88 RMSE 8.7 cm; MERCATOR "
        "r 0.93 RMSE 6.9 cm; mean of the two r 0.94 RMSE 6.5 cm. Without "
        "the inverse barometer MERCATOR scores only r 0.77: it misses the "
        "annual cycle here, which is mostly the air-pressure response. Gauge "
        "residual -0.72 to +0.48 m. " + RESID_METHOD)


def majis_slide(prs):
    """The contrast slide. The surge is made inside the Gulf by wind over a
    shallow enclosed sea; the tide comes in through Hormuz. So outside, at
    Majis, the tide is easy and the surge is small and hard to pick out; at
    the head of the Gulf it is the other way round. Same vertical scale as
    Salmiya, so the difference in size is the first thing seen."""
    return residual_slide(
        prs, "Majis", "Majis, Oman: outside the Gulf,\nthe surge is small.",
        [("±7 cm", "typical surge\nSalmiya ±18 cm"),
         ("0.98", "tide correlation\nSalmiya 0.87"),
         ("0.74", "surge correlation\nSalmiya 0.94")],
        "The surge is generated inside the Gulf by wind over a shallow, "
        "enclosed sea; the tide comes in from outside. At Majis the tide is "
        "near exact and the surge small and hard to pick out; at the head of "
        "the Gulf, the reverse. Gauge record 6 Feb 2023 – 24 Mar 2024.",
        "Majis: 8 637 hourly pairs. CROCO r 0.60 RMSE 6.0 cm; MERCATOR "
        "r 0.77 RMSE 4.3 cm; mean r 0.74 RMSE 4.6 cm - here MERCATOR is the "
        "stronger model and the ensemble mean sits just below it. Residual std: gauge "
        "6.7 cm, CROCO 6.4 cm; at Salmiya gauge 18.4 cm. The spikes are in the "
        "gauge record; removing them lifts r only slightly. " + RESID_METHOD)


# ------------------------------------------------------------------ 03
def abudhabi_slide(prs):
    s = new_slide(prs)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Section 03  ·  Khalifa Port")
    desal.headline(s, MARGIN, Inches(1.30), Inches(11.4),
                   "Nine years of surge at Khalifa Port.", size=34)
    dk.picture(s, asset("wl-abudhabi.png"), MARGIN, Inches(2.08),
               width=Inches(11.0))
    desal.caption(s, MARGIN, Inches(6.88), Inches(11.5),
                  "model only — there is no gauge in this comparison yet. A "
                  "year of your own record turns this into the slides before it.")
    notes(s, "CROCO C04_I01 hindcast, 2016-2024 hourly, at the nearest wet "
             "cell to the terminal (~3 km grid, ~6 m deep, within 2 km). "
             "Residual extracted exactly as tidal_analysis.py does it. "
             "-0.37 to +0.60 m; 0.1 / 99.9 percentiles -0.27 / +0.47 m. "
             "Set-downs below -0.25 m: ~4 events a year; below -0.20 m: ~11 a "
             "year. Zayed Port, 40 km along the coast, tracks it at r 0.99.")
    return s


def tidetable_slide(prs):
    s = new_slide(prs)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Section 03  ·  Khalifa Port")
    desal.headline(s, MARGIN, Inches(1.30), Inches(5.6),
                   "What the tide table\ndidn't say.", size=36)
    desal.bullets(s, MARGIN, Inches(2.95), Inches(5.2), [
        "January 2023, replayed from the hindcast: a week of spring lows with "
        "the water held 0.1 – 0.2 m under the tide throughout.",
        "On 22 January it fell to −1.17 m — the lowest water in nine years, "
        "and 0.11 m below the lowest tide the table predicts in all of them.",
        "Low water was below −1.0 m on five days running. The tide table "
        "says three.",
        "The berth works to the total water level. The tide table gives you "
        "only part of it.",
    ], size=14)
    dk.picture(s, asset("wl-tidetable-khalifa.png"), Inches(6.55),
               Inches(2.25), width=Inches(6.2))
    desal.caption(s, MARGIN, Inches(6.86), Inches(11.5),
                  "hindcast, not a forecast issued at the time — Khalifa Port, "
                  "nearest model cell, levels about model mean sea level")
    notes(s, "18-27 Jan 2023. Lowest total water in 2016-2024: -1.17 m at "
             "17:00 on 22 Jan (tide -1.06, residual -0.11). The lowest "
             "predicted tide anywhere in 2016-2024 is -1.06 m - a stand-in "
             "for the floor of a tide table, not a formal LAT. 25 Jan: tide "
             "-0.89, residual -0.20, water -1.09. Levels are about model MSL; "
             "on the client's chart datum they would shift by a constant.")
    return s


# ------------------------------------------------------------------ kept
def retext_kept(prs, k):
    """k maps the base deck's 1-based slide numbers to slides."""
    # 1. title
    s = k[1]
    set_paras(shape(s, 163), [[("Water levels for your port:\n"
                                "past, present and future", 1)]])
    set_paras(shape(s, 164), [[("PITCH DECK", 0)],
                              [(f"Prepared for [Client]  ·  "
                                f"{date.today():%B %Y}", 0)]])

    # 2. contents
    set_paras(shape(k[2], 173), [
        # the number is split over two mono runs on items 02-04, so the text
        # borrows the paragraph's last run, which is always the body face
        [(f"{i + 1:02d}", 0), (f"    {t}", -1)] for i, t in enumerate([
            "How the model works",
            "Tested against tide gauges",
            "Water levels at your port",
            "What we’d build for you"])])

    # 4-7. the chain slides: only the rail's last stop changes
    for n, sid in ((4, None), (5, 238), (6, None), (7, None)):
        hits = [sh for sh in k[n].shapes if sh.has_text_frame
                and sh.text_frame.text.strip() == "TURBIDITY"]
        if len(hits) != 1:
            raise SystemExit(f"slide {n}: expected one TURBIDITY rail label")
        set_paras(hits[0], [[("WATER LEVEL", 0)]])

    # 12. section divider
    retext(k[12], 373, "SECTION 04")

    # 13. three offerings
    s = k[13]
    retext(s, 381, "SECTION 04")
    for head, body, h, b in (
            (386, 387, "Design hindcast",
             "Ten years of hourly water level, surge, waves and currents at "
             "your port — the design record for a channel, a berth or an "
             "expansion."),
            (391, 392, "A daily UKC forecast",
             "Tide, surge and waves, every morning, seven days ahead, at your "
             "berths and channel — in a dashboard built around your own "
             "clearance limits."),
            (396, 397, "Future climate outlook",
             "The same chain, forced by climate projections and sea level "
             "rise, to ask what your design water levels look like in the "
             "2050s.")):
        retext(s, head, h)
        retext(s, body, b)

    # 14. offering 01
    s = k[14]
    retext(s, 405, "Design hindcast for port expansion.")
    set_paras(shape(s, 406), [dash_item(t) for t in [
        "Ten years of hourly water level, residual, waves and currents at "
        "your channel and berths, from the modelling chain you have just seen.",
        "Extreme water levels for design, high and low, with return periods — "
        "fitted to a decade of model record, not a year of gauge data.",
        "Accessibility statistics: how often, and for how long, a given "
        "draught could not have sailed on a given tide."]])
    set_paras(shape(s, 409), [spec_item(*r) for r in [
        ("period", "2015 → present"),
        ("output", "hourly"),
        ("variables", "WL, residual, Hs, Tp, u, v"),
        ("resolution", "~3 km, nested to ~1 km or finer"),
        ("deliverable", "technical report + the data"),
        ("", "extraction at your berths and channel")]])
    retext(s, 413, "At ~3 km the model resolves the Gulf and your coastline, "
                   "not your basin. Harbour resonance and berth-scale currents "
                   "need a finer nest, which we scope in rather than assume.")

    # 15. dashboard: new words, and a ports view in place of the desal one
    s = k[15]
    set_paras(shape(s, 422), [dash_item(t) for t in [
        "Your berths and channel, on your chart datum",
        "Total water level: tide, surge and waves",
        "Your limits on total water level at each berth, in the units "
        "your pilots use.",
        "Automated warnings (e.g. email) when forecast water level "
        "crosses your limit"]])
    old = shape(s, 423)
    left, top, width = old.left, old.top, old.width
    old._element.getparent().remove(old._element)
    pic = dk.picture(s, asset("wl-dashboard.png"), left, top, width=width)
    # behind the label, which sat over the old picture
    s.shapes._spTree.remove(pic._element)
    shape(s, 424)._element.addprevious(pic._element)
    # the label sat on the old screenshot's quiet corner; this figure has
    # none, so it goes underneath
    lab = shape(s, 424)
    set_paras(lab, [[("ILLUSTRATIVE VIEW  ·  JAN 2023 HINDCAST, REPLAYED", 0)]])
    lab.left, lab.top = pic.left, pic.top + pic.height + Inches(0.12)
    lab.width, lab.height = pic.width, Inches(0.3)

    # 16. with your data
    s = k[16]
    set_paras(shape(s, 433), [dash_item(t) for t in [
        "The model forecasts your port from day one, with nothing required "
        "from you. What follows is an improvement, not a condition.",
        "Your tide gauges are the best training data there is. We train a "
        "correction on the gap between what the model said and what the "
        "gauge measured.",
        "It learns your port's own systematic offsets — the timing of the "
        "tide at your berth, your local bathymetry, your datum.",
        "And it compounds — every month of operation is more training data. "
        "The forecast in year two is not the one you started with."]])
    set_paras(shape(s, 436), [spec_item(*r) for r in [
        ("variable", "water level, any gauge in the port"),
        ("datum", "your chart datum, as surveyed"),
        ("record", "1 year minimum  ·  2+ preferred"),
        ("sampling", "hourly or better  ·  gaps fine"),
        ("method", "learned residual correction, re-fit as data arrives"),
        ("your data", "stays yours, never pooled")]])

    # 17. climate
    s = k[17]
    retext(s, 444, "Future climate outlook - design\n"
                   "water levels in 2050 and beyond.")
    set_paras(shape(s, 445), [dash_item(t) for t in [
        "A quay built today will still be working in the 2070s, but it is "
        "being designed against the water levels of the last few decades.",
        "Sea level is rising, and surge and waves ride on top of it — the "
        "design high water moves with all three.",
        "The same model workflow can be run for future climate scenarios, "
        "with the projected climate and sea level rise built into its "
        "forcing."]])
    set_paras(shape(s, 448), [spec_item(*r) for r in [
        ("method", "dynamical downscaling"),
        ("forcing", "CMIP6, multiple scenarios"),
        ("sea level", "projected rise, built into the model"),
        ("horizon", "2050s and 2080s slices (and others as needed)"),
        ("answers", "extreme water levels, surge and waves"),
        ("built on", "the identical CROCO and WW3"),
        ("", "models used for previous offerings")]])


# ------------------------------------------------------------------ build
KEEP = [1, 2, 3, 4, 5, 6, 7, 12, 13, 14, 15, 16, 17, 18]


def build(out="oma-pitch-ports.pptx"):
    prs = Presentation(BASE)
    base = list(prs.slides)
    if len(base) != 18:
        raise SystemExit(f"{BASE} has {len(base)} slides, expected 18")
    k = {n: base[n - 1] for n in KEEP}
    for n, s in enumerate(base, 1):
        if n not in KEEP:
            drop_slide(prs, s)
    retext_kept(prs, k)

    d02 = divider(prs, "Section 02", "Tested against tide gauges")
    ukc = ukc_slide(prs)
    tide = tide_slide(prs)
    sal = salmiya_slide(prs)
    maj = majis_slide(prs)
    d03 = divider(prs, "Section 03", "Water levels at your port")
    ad = abudhabi_slide(prs)
    tt = tidetable_slide(prs)

    reorder(prs, [k[1], k[2],
                  k[3], k[4], k[5], k[6], k[7],              # 01 the model
                  d02, ukc, tide, sal, maj,                   # 02 the gauges
                  d03, ad, tt,                                # 03 your port
                  k[12], k[13], k[14], k[15], k[16], k[17],   # 04 the offer
                  k[18]])

    # python-pptx names a new slide slide<count+1>.xml, and after four drops
    # that is a name a kept slide still owns -- the zip then carries two parts
    # of one name and PowerPoint refuses the file. Renumber everything in final
    # order; the names are only read at save.
    from pptx.opc.packuri import PackURI
    for i, s in enumerate(prs.slides, 1):
        s.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
        if s.has_notes_slide:
            s.notes_slide.part.partname = PackURI(
                f"/ppt/notesSlides/notesSlide{i}.xml")

    path = os.path.join(dk.HERE, out)
    prs.save(path)
    chain.check_coordinates(path)
    print(f"wrote {path} — {len(prs.slides._sldIdLst)} slides")
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("-o", "--out", default="oma-pitch-ports.pptx")
    build(ap.parse_args().out)
