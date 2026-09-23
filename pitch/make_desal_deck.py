#!/usr/bin/env python3
"""
Ocean Motion Analytics — desalination pitch deck.

Takes the editorial treatments build_template.py established and the ten-slide
"how the models are built" module make_chain_deck.py assembles, and puts them in
one deck aimed at a Gulf desalination operator: what the intake sees, how the
models that predict it are built, and what we would actually sell.

The chain slides are imported rather than copied. make_chain_deck's builders all
take the Presentation as their first argument, so they drop into any deck; the
only thing this file owns is the framing around them. Reword the chain here and
the standalone module goes stale, so don't -- edit make_chain_deck.py and both
decks move together.

Positioning follows oceanmotion-company/strategy/business-model.md, which is the
authority on what may be claimed:
  - lead on turbidity; T and S are context, not the pitch
  - the ladder is a sequence, not a priced menu -- no numbers on these slides
  - HAB / chlorophyll is out of scope and is named as out of scope
  - the daily service is best-effort today; no SLA is implied anywhere

Runs in the base environment (needs python-pptx):
    ~/mambaforge/bin/python3 make_desal_deck.py
Output: oma-pitch-desal.pptx
"""

import argparse
import os
from datetime import date

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import deckkit as dk
from deckkit import (ABYSS, ACCENT, CYAN, CYAN_LT, INK, INK_2, INK_3, NAVY,
                     NAVY_MID, PAPER, PAPER_2, RULE,
                     F_BODY, F_DISPLAY, F_MONO, EMU_W, EMU_H, MARGIN)
import make_chain_deck as chain

ASSET = dk.ASSET_DIR
HERO_STILL = os.path.join(ASSET, "hero-streamlines.png")
HERO_LOOP = os.path.join(ASSET, "hero-streamlines.gif")

# The one intake the deck argues from. January 2022, C04_I02 hindcast, calib_02
# member 03, at the cell nearest the in-situ NTU logger -- every number below was
# read straight out of the output rather than remembered:
#   depth 6.4 m | background 3.0 NTU
#   wind peak 13.5 m/s   21 Jan 12:00   (site, not domain -- the domain peaks at
#                                        08:00, and pairing that with a site
#                                        turbidity peak is where "sixteen hours"
#                                        came from)
#   Hs peak    1.69 m    21 Jan 13:00
#   bed stress 1.04 N/m2 21 Jan 18:30
#   near-bed   17.8 NTU  21 Jan 21:30
#   surface    12.0 NTU  21 Jan 23:30
#   hours >10 NTU over January: 82 near the bed, 19 at the surface
SITE = "24.37 N  ·  54.05 E  ·  6 m"


def slide(prs, ground=PAPER):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    dk.bg(s, ground)
    return s


def headline(s, left, top, width, lead, tail=None, size=42, color=INK,
             accent=ACCENT, height=Inches(1.9)):
    tf = dk.box(s, left, top, width, height)
    p = tf.paragraphs[0]
    p.line_spacing = 1.03
    dk.run(p, lead, font=F_DISPLAY, size=size, color=color)
    if tail:
        dk.run(p, tail, font=F_DISPLAY, size=size, color=accent, italic=True)
    return tf


def caption(s, left, top, width, text, color=INK_2):
    tf = dk.box(s, left, top, width, Inches(0.6))
    dk.run(tf.paragraphs[0], text, font=F_BODY, size=13, color=color,
           italic=True)
    return tf


def bullets(s, left, top, width, items, size=15, color=INK_2, lead=ACCENT):
    tf = dk.box(s, left, top, width, Inches(3.6))
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(11)
        p.line_spacing = 1.32
        dk.run(p, "—  ", font=F_BODY, size=size, color=lead)
        dk.run(p, t, font=F_BODY, size=size, color=color)
    return tf


def metrics(s, items, top=Inches(3.30), big_color=ACCENT, lbl_color=CYAN_LT):
    """The template's three-up figure row. Numbers only, no sentences."""
    mw = Inches(3.6)
    for i, (big, lbl) in enumerate(items):
        x = MARGIN + i * (mw + Inches(0.2))
        tf = dk.box(s, x, top, mw, Inches(1.3))
        dk.run(tf.paragraphs[0], big, font=F_DISPLAY, size=60, color=big_color)
        tf = dk.box(s, x, top + Inches(1.40), mw, Inches(0.8))
        p = tf.paragraphs[0]
        p.line_spacing = 1.32
        dk.run(p, lbl, font=F_MONO, size=11, color=lbl_color, track=120,
               caps=True)


def cards(s, items, top=Inches(2.98), height=Inches(3.44), n=3):
    """n equal cards across the live width, each with a tag, a head and a body.

    Geometry is tight: at three across a card is 3.6 in wide and its body has
    2.9 in of measure, which is about eight lines of 12.5 pt before it runs out
    of card. Keep the bodies under ~230 characters."""
    gap = Inches(0.35)
    cw = dk._emu((Inches(11.5) - gap * (n - 1)) / n)
    for i, (tag, head, body) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        dk.rect(s, x, top, cw, height, PAPER_2)
        dk.rect(s, x, top, cw, Pt(3), ACCENT)
        inset = Inches(0.34)
        tw = cw - inset * 2
        tf = dk.box(s, x + inset, top + Inches(0.30), tw, Inches(0.4))
        dk.run(tf.paragraphs[0], tag, font=F_MONO, size=10, color=ACCENT,
               track=180, caps=True)
        tf = dk.box(s, x + inset, top + Inches(0.72), tw, Inches(0.9))
        p = tf.paragraphs[0]
        p.line_spacing = 1.05
        dk.run(p, head, font=F_DISPLAY, size=24, color=INK)
        tf = dk.box(s, x + inset, top + Inches(1.48), tw, Inches(1.86))
        p = tf.paragraphs[0]
        p.line_spacing = 1.38
        dk.run(p, body, font=F_BODY, size=12.5, color=INK_2)


def divider(prs, tag, lead, tail, ground=NAVY):
    s = slide(prs, ground)
    dk.eyebrow(s, MARGIN, Inches(3.0), tag, color=CYAN)
    headline(s, MARGIN, Inches(3.5), Inches(11.4), lead, tail, size=58,
             color=PAPER)
    return s


def note(s, left, top, width, head, body, on_dark=False):
    """A boxed aside. Used where the deck concedes something -- the one register
    a pitch deck normally has no shape for, and the reason a technical buyer
    believes the rest of it."""
    ground = NAVY_MID if on_dark else PAPER_2
    body_c = CYAN_LT if on_dark else INK_2
    h = Inches(1.84)
    dk.rect(s, left, top, width, h, ground)
    dk.rect(s, left, top, Pt(3), h, ACCENT)
    inset = Inches(0.34)
    tf = dk.box(s, left + inset, top + Inches(0.26), width - inset * 2,
                Inches(0.34))
    dk.run(tf.paragraphs[0], head, font=F_MONO, size=10, color=ACCENT,
           track=170, caps=True)
    tf = dk.box(s, left + inset, top + Inches(0.66), width - inset * 2,
                Inches(1.06))
    p = tf.paragraphs[0]
    p.line_spacing = 1.38
    dk.run(p, body, font=F_BODY, size=12.5, color=body_c)


# ------------------------------------------------------------------ 1. open
def title_slide(prs, hero=HERO_STILL):
    """Still hero rather than the animated loop: this deck is presented over
    Teams more often than in a room, and slow fine detail is the first thing a
    video codec smears. Swap HERO_LOOP back in for a projector."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    dk.full_bleed(s, hero)
    dk.rect(s, MARGIN, Inches(2.4), Inches(1.6), Pt(2), ACCENT)
    dk.wordmark(s, MARGIN, Inches(0.7), on_dark=True)

    tf = dk.box(s, MARGIN, Inches(2.7), Inches(10.5), Inches(2.8))
    p = tf.paragraphs[0]
    p.line_spacing = 1.02
    dk.run(p, "The water your plant", font=F_DISPLAY, size=58, color=PAPER)
    p2 = tf.add_paragraph()
    p2.line_spacing = 1.02
    dk.run(p2, "drinks, ", font=F_DISPLAY, size=58, color=PAPER)
    dk.run(p2, "forecast", font=F_DISPLAY, size=58, color=ACCENT, italic=True)
    dk.run(p2, ".", font=F_DISPLAY, size=58, color=PAPER)

    tf = dk.box(s, MARGIN, Inches(5.9), Inches(9), Inches(1))
    dk.run(tf.paragraphs[0], "Pitch deck  ·  Desalination", font=F_MONO,
           size=11, color=CYAN_LT, track=200, caps=True)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    dk.run(p2, f"Prepared for [Client]  ·  {date.today():%B %Y}",
           font=F_BODY, size=14, color=INK_3)
    return s


CONTENTS = ["What your intake actually sees",
            "How we model it, end to end",
            "Characterising your site",
            "A forecast that runs every day",
            "The next thirty years"]


def contents_slide(prs):
    s = slide(prs, PAPER)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Contents")
    headline(s, MARGIN, Inches(1.4), Inches(9), "What we'll cover.", size=44)
    tf = dk.box(s, MARGIN, Inches(2.9), Inches(10), Inches(4))
    for i, it in enumerate(CONTENTS):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        dk.run(p, f"{i + 1:02d}", font=F_MONO, size=15, color=ACCENT, track=80)
        dk.run(p, "    " + it, font=F_BODY, size=22, color=INK)
    return s


# ------------------------------------------------- 2. the problem
def variables_slide(prs):
    """Three variables, ranked by how much a forecast of them changes a decision
    at an intake. The ranking, and the refusal to sell on T and S, is straight
    out of business-model.md -- a buyer can get SST and Hs from Copernicus for
    nothing, which is exactly why they are context here and not the pitch."""
    s = slide(prs, PAPER)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Section 01")
    headline(s, MARGIN, Inches(1.32), Inches(11.4),
             "Three variables. One of them\narrives ", "without warning.",
             size=38)
    cards(s, [
        ("The wedge", "Turbidity",
         "Shamal-driven, episodic, and it arrives over hours. It is also what "
         "actually hurts: pretreatment overload, SDI spikes, coagulant dosing, "
         "shortened filter runs, production cuts. Nobody forecasts it for you."),
        ("Context", "Temperature",
         "Real SWRO economics — flux, specific energy, salt and boron passage. "
         "But you already measure it, and tomorrow looks much like today. The "
         "value in temperature is seasonal, which makes it a study rather than "
         "a daily number."),
        ("Context", "Salinity",
         "The slowest of the three, and measured. The salinity question with "
         "teeth is not tomorrow's value — it is recirculation: whether your own "
         "outfall, or the plant next door's, reaches your intake."),
    ])
    caption(s, MARGIN, Inches(6.72), Inches(11.5),
            "we keep temperature and salinity on the dashboard because they are "
            "nearly free once the model runs — but turbidity is the reason to "
            "have one")
    return s


def event_slide(prs):
    """The whole argument as three numbers, from one modelled intake."""
    s = slide(prs, NAVY)
    dk.eyebrow(s, MARGIN, Inches(0.95), "One shamal, one intake", color=CYAN)
    headline(s, MARGIN, Inches(1.40), Inches(11.4),
             "21 January 2022. The wind turned\nat midday. By midnight the "
             "water\nwas ", "brown.", size=36, color=PAPER)
    metrics(s, [("3 → 18", "NTU near the bed,\nagainst a 3 NTU background"),
                ("12 h", "from the wind peak\nto the plume at surface"),
                ("4 ×", "longer above 10 NTU near the bed\nthan at the surface")],
            top=Inches(3.62))
    caption(s, MARGIN, Inches(6.86), Inches(11.5),
            f"modelled at {SITE} — and the satellite that would have told you "
            "about it sees only the surface, the smaller half",
            color=CYAN_LT)
    return s


# ------------------------------------------------- 4. the offering
def offering_slide(prs):
    s = slide(prs, PAPER)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Section 03")
    headline(s, MARGIN, Inches(1.32), Inches(11.4),
             "Three ways in, in the order\nthey usually ", "happen.", size=38)
    cards(s, [
        ("First", "Site characterisation",
         "A hindcast study at your intake. What the water has actually done — "
         "turbidity, temperature, salinity, currents — over ten years, with "
         "every event resolved hour by hour."),
        ("Then", "A daily forecast",
         "The same models, run every morning without anyone touching them, "
         "through a dashboard designed around your intake depth and your own "
         "thresholds. Seven days ahead."),
        ("Later", "Climate outlook",
         "The same chain, forced by climate projections instead of reanalysis, "
         "to ask what the intake looks like in the 2050s. The one thing here we "
         "have not yet delivered."),
    ])
    caption(s, MARGIN, Inches(6.72), Inches(11.5),
            "a sequence, not a menu — each step is built out of what the one "
            "before it produced, which is why the second is cheap once the "
            "first exists")
    return s


def characterisation_slide(prs):
    s = slide(prs, PAPER)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Offering 01")
    headline(s, MARGIN, Inches(1.38), Inches(6.6),
             "Start with what the water\nhas ", "already done.", size=36)
    bullets(s, MARGIN, Inches(3.05), Inches(6.5), [
        "Ten years of hourly conditions at your intake, from the chain you "
        "have just seen — not a nearby buoy, not a regional average.",
        "Every shamal since 2015 resolved: how fast turbidity climbed, how "
        "high, how long it held, at the depth you actually draw from.",
        "The numbers a design or an operating procedure needs: exceedance "
        "statistics, event durations, seasonal envelopes.",
        "Brine recirculation, where that is the question: whether your outfall, "
        "or your neighbour's, reaches your own intake.",
    ])
    chain.specs(s, Inches(8.20), Inches(1.42), Inches(4.6),
                [("period", "2015 → present"),
                 ("output", "hourly  ·  all 15 levels"),
                 ("variables", "T, S, u, v, Hs, Tp, NTU"),
                 ("resolution", "3.5 km, nested to ~1 km"),
                 ("deliverable", "technical report + the data"),
                 ("", "extraction at your own"),
                 ("", "intake coordinates and depth")],
                "what a study covers")
    note(s, Inches(8.20), Inches(4.72), Inches(4.6), "Scope, stated up front",
         "Turbidity here is sediment-driven — shamal-stirred plumes and wadi "
         "input. Algal blooms and biological fouling are not modelled, and we "
         "would rather say so than be asked.")
    caption(s, MARGIN, Inches(6.72), Inches(11.5),
            "the study stands on its own — it is not a sales doorbell for the "
            "subscription, and it is scoped to a question you already have")
    return s


def daily_slide(prs):
    """The operational fact, given its own slide because it is the thing most
    easily lost: this is not a study that could be re-run, it is a system that
    already runs unattended every morning and has been for months."""
    s = slide(prs, NAVY)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Offering 02", color=CYAN)
    headline(s, MARGIN, Inches(1.38), Inches(11.4),
             "These models already run\nevery ", "morning.", size=40,
             color=PAPER)
    metrics(s, [("Daily", "one full cycle, 00Z,\nfired at 06:00 UTC"),
                ("7 days", "forecast horizon,\nhourly throughout"),
                ("75 min", "from raw global data\nto a published forecast")],
            top=Inches(3.20))
    note(s, MARGIN, Inches(5.62), Inches(11.5), "What we will not claim",
         "Today this is best-effort: automated, monitored, and delivered every "
         "morning — but without a contractual uptime guarantee, because one "
         "person cannot honour a 24/7 SLA. We price it as what it is. Early "
         "users are supplementing what they already have, not betting the plant.",
         on_dark=True)
    return s


def dashboard_slide(prs):
    s = slide(prs, ABYSS)
    dk.eyebrow(s, MARGIN, Inches(0.72), "Offering 02", color=CYAN)
    headline(s, MARGIN, Inches(1.12), Inches(7.0),
             "Built with you, not\nhanded ", "to you.", size=34, color=PAPER)
    tf = dk.box(s, MARGIN, Inches(2.60), Inches(5.9), Inches(3.4))
    for i, t in enumerate([
            "Your intake, at your depth — not the nearest grid point at the "
            "surface.",
            "Your thresholds and your alarm levels, in the units your operators "
            "already use.",
            "Turbidity shown as all five ensemble members, so you see the "
            "spread rather than a single number pretending to certainty.",
            "Water level, waves, temperature, salinity and currents alongside "
            "it, because they cost nothing once the model runs."]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        p.line_spacing = 1.36
        dk.run(p, "—  ", font=F_BODY, size=14, color=ACCENT)
        dk.run(p, t, font=F_BODY, size=14, color=CYAN_LT)

    ph = dk.rect(s, Inches(7.30), Inches(1.98), Inches(5.10), Inches(4.06),
                 NAVY_MID)
    ph.line.color.rgb = CYAN
    ph.line.width = Pt(1)
    tf = dk.box(s, Inches(7.30), Inches(3.80), Inches(5.10), Inches(0.6),
                anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    dk.run(p, "[ Replace with dashboard screenshot ]", font=F_MONO, size=11,
           color=CYAN_LT, track=120, caps=True)
    caption(s, MARGIN, Inches(6.42), Inches(11.5),
            "the same public site that runs the Gulf today — a client view is a "
            "point extraction, a threshold config and a login",
            color=INK_3)
    return s


def climate_slide(prs):
    """The offering that does not exist yet, labelled as such on the slide.

    Naming it in-development is not a hedge -- in a market this sceptical it is
    the reason the rest of the deck is believed, and it converts the gap into
    the thing a first client part-funds."""
    s = slide(prs, PAPER)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Offering 03  ·  In development")
    headline(s, MARGIN, Inches(1.38), Inches(7.0),
             "And what the intake\nlooks like ", "in 2050.", size=36)
    bullets(s, MARGIN, Inches(3.05), Inches(6.6), [
        "A plant commissioned today will still be drawing water in the 2050s. "
        "It is being designed against the last ten years.",
        "The Gulf is warming, and its salinity is rising — in part under the "
        "cumulative brine load of the region's own plants.",
        "Both move the numbers a train is sized against: feed temperature sets "
        "flux, energy and boron passage; feed salinity sets recovery.",
        "The chain does not care what forces it — swap reanalysis for "
        "downscaled projections and it answers 2050, not tomorrow.",
    ])
    chain.specs(s, Inches(8.20), Inches(1.42), Inches(4.6),
                [("method", "dynamical downscaling"),
                 ("forcing", "CMIP6, two scenarios"),
                 ("horizon", "2050s and 2080s slices"),
                 ("answers", "T, S and NTU at the intake"),
                 ("built on", "the identical CROCO, WW3"),
                 ("", "and sediment configuration"),
                 ("status", "capability, not yet a product")],
                "how it would work")
    note(s, Inches(8.20), Inches(4.72), Inches(4.6), "Where this actually is",
         "The model chain exists and has been assessed. The climate forcing and "
         "the downscaling protocol have not been built. We would rather name "
         "that than imply a track record we do not have.")
    caption(s, MARGIN, Inches(6.72), Inches(11.5),
            "this is the piece a first client part-funds — and the reason to "
            "have the conversation before the design is frozen")
    return s


# ------------------------------------------------------------------ 5. close
def statement_slide(prs):
    s = slide(prs, PAPER)
    tf = dk.box(s, MARGIN, 0, Inches(11.5), EMU_H, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.06
    dk.run(p, "The people who have to know\nwhat the sea is doing ",
           font=F_DISPLAY, size=46, color=INK)
    dk.run(p, "shouldn't\nbe guessing", font=F_DISPLAY, size=46, color=ACCENT,
           italic=True)
    dk.run(p, ".", font=F_DISPLAY, size=46, color=INK)
    return s


def closing_slide(prs):
    s = slide(prs, ABYSS)
    dk.wordmark(s, MARGIN, Inches(0.7), on_dark=True)
    dk.rect(s, MARGIN, Inches(2.18), Inches(1.6), Pt(2), ACCENT)
    headline(s, MARGIN, Inches(2.48), Inches(11), "Tell us what you're\n",
             "working on.", size=52, color=PAPER)

    # The ask, made explicit. go-to-market.md is unambiguous that the decided
    # ask at this stage is a scoped study, not a subscription.
    tf = dk.box(s, MARGIN, Inches(4.30), Inches(7.4), Inches(0.9))
    p = tf.paragraphs[0]
    p.line_spacing = 1.38
    dk.run(p, "The next step we would propose is a scoped characterisation "
              "study at one intake — yours. Thirty minutes first, to work out "
              "whether it maps to a problem you actually have.",
           font=F_BODY, size=15, color=CYAN_LT)

    tf = dk.box(s, MARGIN, Inches(5.68), Inches(10), Inches(1.2))
    dk.run(tf.paragraphs[0], "Giles Fearon", font=F_BODY, size=18, color=PAPER,
           bold=True)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    dk.run(p2, "giles@oceanmotionanalytics.com", font=F_MONO, size=13,
           color=CYAN_LT)
    p3 = tf.add_paragraph()
    p3.space_before = Pt(2)
    dk.run(p3, "oceanmotionanalytics.com", font=F_MONO, size=13, color=INK_3)
    return s


# ------------------------------------------------------------------ build
def build(out="oma-pitch-desal.pptx"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = EMU_W, EMU_H
    dk.set_theme_fonts(prs)

    title_slide(prs)                                            # 1
    contents_slide(prs)                                         # 2

    divider(prs, "Section 01", "What your intake\n", "actually sees.")  # 3
    variables_slide(prs)                                        # 4
    event_slide(prs)                                            # 5

    divider(prs, "Section 02", "How we ", "model it.", ground=ABYSS)    # 6
    # --- the chain module, imported whole rather than restated
    chain.overview(prs)                                         # 7
    chain.link_slide(prs, chain.SLIDES[0])                      # 8  atmosphere
    chain.boundaries_slide(prs)                                 # 9
    chain.grid_slide(prs)                                       # 10
    chain.circulation_slide(prs)                                # 11
    chain.link_slide(prs, chain.SLIDES[1])                      # 12 waves
    chain.process_slide(prs)                                    # 13
    chain.erosion_slide(prs)                                    # 14
    chain.turbidity_slide(prs)                                  # 15
    if os.path.exists(os.path.join(ASSET, chain.BOOKENDS["closer"]["img"])):
        chain.bookend(prs, chain.BOOKENDS["closer"])            # 16
    else:
        print("  (no cascade GIF — run make_model_chain.py --gif)")

    divider(prs, "Section 03", "What we'd\n", "build for you.")  # 17
    offering_slide(prs)                                         # 18
    characterisation_slide(prs)                                 # 19
    daily_slide(prs)                                            # 20
    dashboard_slide(prs)                                        # 21
    climate_slide(prs)                                          # 22

    statement_slide(prs)                                        # 23
    closing_slide(prs)                                          # 24

    dk.embed_fonts(prs)
    path = os.path.join(dk.HERE, out)
    prs.save(path)
    chain.check_coordinates(path)
    print(f"wrote {path} — {len(prs.slides._sldIdLst)} slides")
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("-o", "--out", default="oma-pitch-desal.pptx")
    build(ap.parse_args().out)
