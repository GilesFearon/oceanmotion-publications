#!/usr/bin/env python3
"""
Ocean Motion Analytics — supplier deck for a UAE coastal engineering
consultancy (ADEC).

The audience is a consultancy, not an asset owner: they run their own
nearshore models and deliver the client's study. So the deck sells what sits
behind their work -- regional design data, turbidity, climate horizons -- and
never port-scale downscaling, which is their business.

It is also a discovery deck. The 40-year design-conditions dataset is built in
H1 2027 (see oceanmotion-company/business-plan/); today there is a 10-year
circulation hindcast and one year of waves. Slides say so. The ask is a
founding-partner commitment -- a letter of intent and validation data -- not
a purchase.

Built the same way as make_ports_deck.py, and reusing its helpers: open the
desal deck as sent (it carries hand edits the generators don't know about),
keep the model and turbidity slides, retext the offering slides as the dataset,
climate and ways-of-working slides, and add the surge evidence and new slides
with deckkit. Shape ids are the base file's; every lookup fails loudly.

Story:
  --  what has changed since we last worked together
  01  how the model works         inputs, grid, circulation, waves, turbidity
  02  tested against tide gauges  Salmiya surge; nine years at Khalifa Port
  03  a design dataset            the dataset (status stated); climate horizons
  04  working together            three ways; founding partner

Prices are deliberately not on the slides or in the notes (a sent .pptx
carries its notes). They live in adec-meeting-prep.md.

Runs in the base environment (needs python-pptx):
    python3 make_adec_deck.py
Output: oma-pitch-adec.pptx
"""

import argparse
import os
from datetime import date

from pptx import Presentation
from pptx.util import Inches

import deckkit as dk
from deckkit import ACCENT, MARGIN
import make_desal_deck as desal
import make_ports_deck as ports
from make_ports_deck import (shape, set_paras, dash_item, spec_item, retext,
                             drop_slide, reorder, new_slide, notes, divider)

BASE = ports.BASE
CLIENT = "ADEC"


# ------------------------------------------------------------------ new
def since_slide(prs):
    s = new_slide(prs)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Since we last worked together")
    desal.headline(s, MARGIN, Inches(1.35), Inches(11.4),
                   "A model of the whole Gulf, ", "running every day.", size=40)
    desal.cards(s, [
        ("Operational", "A daily Gulf forecast",
         "Circulation, water level and waves across the whole Gulf, seven "
         "days ahead, every day, on a fully automated chain."),
        ("Tested", "A decade of hindcast",
         "2015 to 2025, hourly, checked against tide gauges, satellite "
         "temperature and satellite wave heights. The surge is the standout."),
        ("Unique", "Sediment turbidity",
         "Shamal-driven turbidity from waves and currents at the bed, tuned "
         "to in-situ measurements. We know of no one else modelling it for "
         "the Gulf."),
    ], top=Inches(2.85))
    notes(s, "Open with the relationship, then this slide. The point: the "
             "regional model they would otherwise have to build or buy "
             "already exists, runs daily, and is tested. Everything later in "
             "the deck is built on it.")
    return s


def evidence_rail(s, active):
    """The validation slides sit outside the chain's progress rail, so they
    carry a section eyebrow where the rail would be."""
    dk.eyebrow(s, MARGIN, Inches(0.42), "Section 02  ·  Tested against tide "
                                        "gauges")


def salmiya_slide(prs):
    """The surge evidence with CROCO alone. The ports deck draws MERCATOR
    beside it to argue for the operational ensemble; for design data the
    argument is the model's own long record, and MERCATOR's long-record
    counterpart (GLORYS) keeps only daily means, so it is left off."""
    s = new_slide(prs)
    evidence_rail(s, None)
    desal.headline(s, MARGIN, Inches(1.00), Inches(5.6),
                   "Salmiya, Kuwait: the surge,\nevent by event.", size=30)
    for k, (big, lbl) in enumerate([
            ("0.88", "correlation\nhourly, 9 months"),
            ("8.7 cm", "RMSE"),
            ("±18 cm", "typical surge\nat the gauge")]):
        ports.stat(s, Inches(6.75) + k * Inches(2.0), Inches(0.95), big, lbl,
                   width=Inches(1.9), size=40)
    dk.picture(s, ports.asset("wl-residual-salmiya-croco.png"), MARGIN,
               Inches(2.62), width=Inches(11.0))
    desal.caption(s, MARGIN, Inches(6.40), Inches(11.5),
                  "non-tidal residual only: gauge and model put through the "
                  "same harmonic analysis over the gauge record, "
                  "15 Jun 2023 – 26 Mar 2024")
    notes(s, "Salmiya: 6 834 hourly pairs (6 gauge spikes removed). CROCO "
             "r 0.88, RMSE 8.7 cm. Gauge residual -0.72 to +0.48 m, std "
             "18.4 cm. Residual = series minus its own utide fit over the "
             "gauge record period, trend=False. The surge is generated inside "
             "the Gulf by wind over a shallow enclosed sea, so it is largest "
             "at the head of the Gulf and along the southern shelf.")
    return s


def abudhabi_slide(prs):
    s = new_slide(prs)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Section 02  ·  Abu Dhabi")
    desal.headline(s, MARGIN, Inches(1.30), Inches(11.4),
                   "Nine years of surge at Khalifa Port.", size=34)
    dk.picture(s, ports.asset("wl-abudhabi.png"), MARGIN, Inches(2.08),
               width=Inches(11.0))
    desal.caption(s, MARGIN, Inches(6.88), Inches(11.5),
                  "model only — there is no UAE gauge in this comparison yet. "
                  "Your records are what would change that.")
    notes(s, "CROCO C04_I01 hindcast, 2016-2024 hourly, nearest wet cell to "
             "Khalifa Port (~3 km grid, ~6 m deep). Residual -0.37 to +0.60 m; "
             "0.1 / 99.9 percentiles -0.27 / +0.47 m. Set-downs below -0.25 m "
             "about four a year. Zayed Port, 40 km along the coast, tracks it "
             "at r 0.99. This is the lead-in to the data-sharing ask: the gap "
             "in the evidence is a UAE gauge, and they may hold one.")
    return s


def partner_slide(prs):
    s = new_slide(prs)
    dk.eyebrow(s, MARGIN, Inches(0.95), "Section 04")
    desal.headline(s, MARGIN, Inches(1.30), Inches(11.4),
                   "Founding partner, ", "2027.", size=40)
    col_w = Inches(5.5)
    right = MARGIN + col_w + Inches(0.5)
    dk.eyebrow(s, MARGIN, Inches(2.55), f"What {CLIENT} gets", color=ACCENT)
    desal.bullets(s, MARGIN, Inches(3.00), col_w, [
        "A first-year regional licence at a founding rate.",
        "A say in what the dataset covers: your sites, your variables, your "
        "design horizons.",
        "Early access to each release as it is validated.",
        "Validation at your own stations, published as ours to stand behind.",
    ], size=13)
    dk.eyebrow(s, right, Inches(2.55), "What we'd ask", color=ACCENT)
    desal.bullets(s, right, Inches(3.00), col_w, [
        "A letter of intent: the licence, if the validation clears the bar "
        "we agree now.",
        "Access to gauge, ADCP or buoy records for validating extremes. They "
        "stay yours and are never redistributed.",
        "A reference, once delivered.",
    ], size=13)
    desal.note(s, MARGIN, Inches(5.52), Inches(11.5), "Where it stands",
               "Today: a tested 10-year circulation hindcast and one year of "
               "waves. In build, first half of 2027: the 40-year record, "
               "validated extremes, turbidity climatology and climate "
               "horizons. Founding partners shape it before it is fixed.")
    notes(s, "The ask. A letter of intent conditional on validation is easy "
             "to sign and is the evidence the funding case needs. The data "
             "access matters as much as the money: UAE gauge extremes are the "
             "missing piece of the validation.")
    return s


# ------------------------------------------------------------------ kept
def retext_kept(k):
    """k maps the base deck's 1-based slide numbers to slides."""
    # 1. title
    s = k[1]
    set_paras(shape(s, 163), [[("Regional ocean data\n"
                                "for your Gulf projects", 1)]])
    set_paras(shape(s, 164), [[("PITCH DECK", 0)],
                              [(f"Prepared for {CLIENT}  ·  "
                                f"{date.today():%B %Y}", 0)]])

    # 2. contents
    set_paras(shape(k[2], 173), [
        [(f"{i + 1:02d}", 0), (f"    {t}", -1)] for i, t in enumerate([
            "How the model works",
            "Tested against tide gauges",
            "A design dataset for the Gulf",
            "Working together"])])

    # 8, 10, 11. turbidity: same slides, the consultancy's reasons in notes
    notes(k[8], "For a consultancy the turbidity story is dredging and "
                "reclamation compliance and EIA baselines as much as intakes: "
                "what background turbidity looks like through a shamal "
                "season, and how often a limit is exceeded with no works at "
                "all.")
    notes(k[10], "Calibrated at one station over one month; a multi-year "
                 "climatology validated against satellite turbidity is part "
                 "of the 2027 build. Say so if asked.")

    # 13. three offerings -> three ways to work together
    s = k[13]
    retext(s, 381, "SECTION 04")
    set_paras(shape(s, 382), [[("Three ways to ", 0), ("work together.", 1)]])
    for tag, head, body, t, h, b in (
            (385, 386, 387, "PER PROJECT", "Project data",
             "Design conditions and boundary data at the offshore edge of "
             "your project, with a skill appendix, delivered in days for "
             "your own nearshore models."),
            (390, 391, 392, "ANNUAL", "A regional licence",
             "Unlimited point extractions across the Gulf for your projects "
             "for a year, under one agreement. Your team pulls what each "
             "study needs."),
            (395, 396, 397, "SUBCONTRACT", "Specialist modelling",
             "What most consultancies don't run in-house: sediment turbidity "
             "and dredge plumes, and climate horizons for a design basis.")):
        retext(s, tag, t)
        retext(s, head, h)
        retext(s, body, b)
    notes(s, "We supply what sits behind your study; the nearshore modelling "
             "and design stay with you.")

    # 14. offering 01 -> the dataset
    s = k[14]
    retext(s, 404, "SECTION 03")
    retext(s, 405, "A design dataset for the whole Gulf.")
    set_paras(shape(s, 406), [dash_item(t) for t in [
        "Forty-plus years of hourly waves, water level, surge and currents, "
        "from the modelling chain you have just seen.",
        "Extremes checked against tide gauges and satellite wave heights, "
        "with the skill published alongside the data.",
        "A turbidity climatology: how often, how high and for how long, "
        "at any site and depth.",
        "Climate horizons at 2050 and 2100 for the same variables."]])
    retext(s, 408, "WHAT IT WILL COVER")
    set_paras(shape(s, 409), [spec_item(*r) for r in [
        ("period", "40+ years, hourly"),
        ("variables", "waves, water level, surge"),
        ("", "currents, T, S, turbidity"),
        ("resolution", "~3 km regional"),
        ("deliverable", "point data + skill appendix"),
        ("", "extracted at any coordinates, in days")]])
    retext(s, 412, "STATUS AND SCOPE")
    retext(s, 413, "In build, first half of 2027; today's record is ten years "
                   "of circulation and one of waves. Regional, not nearshore: "
                   "transformation to your structures stays in your models.")
    notes(s, "Be exact about status. Today: 10-year CROCO hindcast (2015-2025) "
             "and one year of coupled waves (2016). The 40-year record, "
             "extremes validation, turbidity climatology and climate horizons "
             "are the 2027 build.")

    # 17. climate
    s = k[17]
    retext(s, 443, "SECTION 03")
    retext(s, 444, "Design conditions at 2050 and 2100,\n"
                   "ready to cite.")
    set_paras(shape(s, 445), [dash_item(t) for t in [
        "Structures designed now will work through the 2070s, but climate "
        "allowances in the Gulf are mostly desk estimates.",
        "Sea level rise lets larger waves reach shallow structures, and "
        "surge rides on a higher mean — the design values move together.",
        "The same rise reduces bed stress, so turbidity can fall at deeper "
        "sites while waves at the structure grow. One consistent set of "
        "runs answers both."]])
    set_paras(shape(s, 448), [spec_item(*r) for r in [
        ("method", "delta (pseudo-global-warming) downscaling"),
        ("scenarios", "SSP2-4.5 and SSP5-8.5"),
        ("horizons", "2050 and 2100"),
        ("sea level", "projected rise, built into the model"),
        ("answers", "design water level, Hs, turbidity"),
        ("built on", "the same validated models")]])
    notes(s, "The method keeps the observed sequence of shamals and adds the "
             "projected change on top; wind changes are handled as a response "
             "curve rather than a claim about how shamals will change.")


# ------------------------------------------------------------------ build
KEEP = [1, 2, 3, 4, 5, 6, 7, 8, 10, 11, 13, 14, 17, 18]


def build(out="oma-pitch-adec.pptx"):
    prs = Presentation(BASE)
    base = list(prs.slides)
    if len(base) != 18:
        raise SystemExit(f"{BASE} has {len(base)} slides, expected 18")
    k = {n: base[n - 1] for n in KEEP}
    for n, s in enumerate(base, 1):
        if n not in KEEP:
            drop_slide(prs, s)
    retext_kept(k)

    since = since_slide(prs)
    d02 = divider(prs, "Section 02", "Tested against tide gauges")
    sal = salmiya_slide(prs)
    ad = abudhabi_slide(prs)
    d03 = divider(prs, "Section 03", "A design dataset for the Gulf")
    d04 = divider(prs, "Section 04", "Working together")
    partner = partner_slide(prs)

    reorder(prs, [k[1], since, k[2],
                  k[3], k[4], k[5], k[6], k[7], k[8], k[10], k[11],  # 01
                  d02, sal, ad,                                       # 02
                  d03, k[14], k[17],                                  # 03
                  d04, k[13], partner,                                # 04
                  k[18]])

    from pptx.opc.packuri import PackURI
    for i, s in enumerate(prs.slides, 1):
        s.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
        if s.has_notes_slide:
            s.notes_slide.part.partname = PackURI(
                f"/ppt/notesSlides/notesSlide{i}.xml")

    path = os.path.join(dk.HERE, out)
    prs.save(path)
    ports.chain.check_coordinates(path)
    print(f"wrote {path} — {len(prs.slides._sldIdLst)} slides")
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("-o", "--out", default="oma-pitch-adec.pptx")
    build(ap.parse_args().out)
