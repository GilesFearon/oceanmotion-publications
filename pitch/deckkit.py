#!/usr/bin/env python3
"""
Ocean Motion Analytics — shared PowerPoint style kit.

The tokens, font machinery and shape helpers that build_template.py established,
lifted out so more than one deck script can use them. Everything here takes the
Presentation (or a slide) as an argument rather than reaching for module state,
which is the one thing that stops build_template.py being importable: it builds
and saves the template at import time.

build_template.py is deliberately left alone. When it next needs editing it can
be switched over to these helpers; until then the duplication is the cheaper
side of the trade.
"""

import os

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ------------------------------------------------------------------ tokens
ABYSS    = RGBColor(0x05, 0x0f, 0x1c)
NAVY     = RGBColor(0x09, 0x1c, 0x33)
NAVY_MID = RGBColor(0x13, 0x29, 0x4a)
CYAN     = RGBColor(0x4f, 0xc3, 0xd7)
CYAN_LT  = RGBColor(0xa8, 0xe0, 0xea)
ACCENT   = RGBColor(0xff, 0x6a, 0x2b)
PAPER    = RGBColor(0xf4, 0xef, 0xe3)
PAPER_2  = RGBColor(0xef, 0xe9, 0xd9)
INK      = RGBColor(0x0b, 0x14, 0x20)
INK_2    = RGBColor(0x4a, 0x56, 0x65)
INK_3    = RGBColor(0x8a, 0x94, 0xa2)
RULE     = RGBColor(0xd6, 0xcd, 0xb6)

F_DISPLAY = "Instrument Serif"
F_BODY    = "IBM Plex Sans"
F_MONO    = "IBM Plex Mono"

HERE      = os.path.dirname(os.path.abspath(__file__))
ASSET_DIR = os.path.join(HERE, "assets")
FONT_DIR  = os.path.join(HERE, "fonts")
MARK_DARK  = os.path.join(ASSET_DIR, "om-mark-dark.png")
MARK_LIGHT = os.path.join(ASSET_DIR, "om-mark-light.png")

EMBED_FONTS = {
    F_DISPLAY: {"regular": "InstrumentSerif-Regular.ttf",
                "italic":  "InstrumentSerif-Italic.ttf"},
    F_BODY:    {"regular": "IBMPlexSans-Regular.ttf",
                "bold":    "IBMPlexSans-Bold.ttf",
                "italic":  "IBMPlexSans-Italic.ttf",
                "boldItalic": "IBMPlexSans-BoldItalic.ttf"},
    F_MONO:    {"regular": "IBMPlexMono-Regular.ttf"},
}

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.92)


# ------------------------------------------------------------------ fonts
def set_theme_fonts(prs, major=F_DISPLAY, minor=F_BODY):
    """Point the theme's font scheme at the brand faces so any text box added
    later inherits them instead of Calibri."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    theme = etree.fromstring(part.blob)
    scheme = theme.find(qn('a:themeElements')).find(qn('a:fontScheme'))
    for tag, face in ((qn('a:majorFont'), major), (qn('a:minorFont'), minor)):
        scheme.find(tag).find(qn('a:latin')).set('typeface', face)
    part._blob = etree.tostring(theme, xml_declaration=True,
                                encoding='UTF-8', standalone=True)


def embed_fonts(prs, families=None):
    """Embed the brand faces as /ppt/fonts/*.fntdata. Naming a font in a run
    states a preference only; a machine without it substitutes silently while
    still reporting the requested name. Honoured by PowerPoint and LibreOffice."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI

    families = EMBED_FONTS if families is None else families
    pres_part = prs.part
    root = pres_part._element
    lst = root.makeelement(qn('p:embeddedFontLst'), {})

    n = 0
    for typeface, faces in families.items():
        ef = root.makeelement(qn('p:embeddedFont'), {})
        ef.append(root.makeelement(qn('p:font'),
                                   {'typeface': typeface, 'charset': '0'}))
        for style in ("regular", "bold", "italic", "boldItalic"):
            if style not in faces:
                continue
            n += 1
            with open(os.path.join(FONT_DIR, faces[style]), "rb") as fh:
                blob = fh.read()
            part = Part(PackURI("/ppt/fonts/font%d.fntdata" % n),
                        "application/x-fontdata", pres_part.package, blob)
            rId = pres_part.relate_to(part, RT.FONT)
            ef.append(root.makeelement(qn("p:" + style), {qn('r:id'): rId}))
        lst.append(ef)

    root.find(qn('p:notesSz')).addnext(lst)


# ------------------------------------------------------------------ shapes
def _emu(v):
    """Coerce a coordinate to whole EMU.

    OOXML types every offset and extent as xsd:long, and python-pptx writes
    whatever it is handed. A stray Python float -- `Inches(2) / 2` is a float,
    not an Emu -- lands in the XML as x="5463539.5", which LibreOffice tolerates
    and PowerPoint and Google Slides both reject outright, with no indication of
    which element is at fault. Every helper here rounds, so arithmetic on
    positions cannot silently produce a file that will not open."""
    return int(round(v))


def flatten(shape):
    """Drop the autoshape's themed <p:style> (it carries an effectRef shadow) so
    the shape renders perfectly flat, matching the website."""
    el = shape._element
    style = el.find(qn('p:style'))
    if style is not None:
        el.remove(style)
    shape.shadow.inherit = False
    return shape


def _to_back(s, shape):
    s.shapes._spTree.remove(shape._element)
    s.shapes._spTree.insert(2, shape._element)
    return shape


def bg(s, color):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    return _to_back(s, flatten(r))


def full_bleed(s, path):
    """Cover the slide with an image, pushed to the back. The source must be 16:9
    or it will be stretched -- there is no cropping here."""
    pic = s.shapes.add_picture(path, 0, 0, width=EMU_W, height=EMU_H)
    return _to_back(s, pic)


def picture(s, path, left, top, **kw):
    kw = {k: (_emu(v) if k in ("width", "height") else v)
          for k, v in kw.items()}
    return s.shapes.add_picture(path, _emu(left), _emu(top), **kw)


def rect(s, left, top, w, h, color, line=None):
    left, top, w, h = (_emu(v) for v in (left, top, w, h))
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line; r.line.width = Pt(1)
    return flatten(r)


def spacing(run_, centipoints):
    """Letter-spacing via the a:rPr spc attribute (1/100 pt)."""
    run_.font._rPr.set('spc', str(int(centipoints)))


def box(s, left, top, w, h, anchor=MSO_ANCHOR.TOP):
    left, top, w, h = (_emu(v) for v in (left, top, w, h))
    tb = s.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def run(p, text, font=F_BODY, size=18, color=INK, bold=False, italic=False,
        track=None, caps=False):
    r = p.add_run()
    r.text = text.upper() if caps else text
    f = r.font
    f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color
    if track is not None:
        spacing(r, track)
    return r


def eyebrow(s, left, top, text, color=INK_2, w=Inches(6)):
    """Mono uppercase eyebrow with a short leading rule (mirrors .section-eyebrow)."""
    rule_w = Inches(0.42)
    rect(s, left, top + Pt(7), rule_w, Pt(1.2), color)
    tf = box(s, left + rule_w + Inches(0.16), top, w, Inches(0.4))
    run(tf.paragraphs[0], text, font=F_MONO, size=12, color=color,
        track=180, caps=True)
    return tf


def glyph(s, left, top, height, on_dark=True):
    return s.shapes.add_picture(MARK_DARK if on_dark else MARK_LIGHT,
                                left, top, height=height)


def arrow(s, x0, y0, x1, y1, color=INK_3, width=Pt(1.25), head="triangle"):
    """Straight connector with an arrowhead.

    python-pptx exposes connectors but not line-end decoration, so the head is
    written onto the line's <a:ln> directly. Used by the chain schematic, which
    is drawn as native shapes rather than as one flat image so the boxes and
    arrows can be nudged in PowerPoint."""
    from pptx.enum.shapes import MSO_CONNECTOR
    x0, y0, x1, y1 = (_emu(v) for v in (x0, y0, x1, y1))
    cxn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x0, y0, x1, y1)
    cxn.line.color.rgb = color
    cxn.line.width = width
    ln = cxn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'),
                          {'type': head, 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return cxn


def node_box(s, left, top, w, h, img=None, label=None, sub=None,
             edge=RULE, label_size=11):
    left, top, w, h = (_emu(v) for v in (left, top, w, h))
    """A labelled box for the chain schematic: optional thumbnail, name above,
    caption below. Every part is a separate shape, so the whole diagram stays
    editable on the slide."""
    r = rect(s, left, top, w, h, PAPER, line=edge)
    if img:
        pad = Pt(4)
        s.shapes.add_picture(img, left + pad, top + pad,
                             width=w - 2 * pad, height=h - 2 * pad)
    if label:
        tf = box(s, left, top - Inches(0.30), w + Inches(1.2), Inches(0.28))
        run(tf.paragraphs[0], label, font=F_BODY, size=label_size, color=INK)
    if sub:
        tf = box(s, left, top + h + Inches(0.06), w + Inches(1.2), Inches(0.26))
        run(tf.paragraphs[0], sub, font=F_MONO, size=8.5, color=INK_3)
    return r


def gradient_text(p, text, c_from, c_to, font, size):
    """Stand-in for the website's `linear-gradient(90deg, ...)` + `background-clip:
    text` on .wordmark__name, by interpolating a solid colour per character.

    PowerPoint can do a true gradient text fill, but Google Slides drops it on
    import and flattens the text to one colour -- so solid runs are used instead.
    They survive every renderer. The cost is that colour steps with character
    index rather than x-position, and kerning is not applied between runs; at
    wordmark size neither is perceptible."""
    n = len(text)
    for i, ch in enumerate(text):
        t = i / (n - 1) if n > 1 else 0.0
        color = RGBColor(*(round(a + (b - a) * t) for a, b in zip(c_from, c_to)))
        run(p, ch, font=font, size=size, color=color)


def wordmark(s, left, top, on_dark=True):
    """The site header lockup: eddy mark, then name + suffix."""
    g_h = Inches(0.46)
    glyph(s, _emu(left - Inches(0.05)), _emu(top - Inches(0.03)), g_h, on_dark)
    tf = box(s, left + g_h, top, Inches(6), Inches(0.5))
    p = tf.paragraphs[0]
    grad = (PAPER, CYAN_LT) if on_dark else (INK, NAVY_MID)
    suf_c = CYAN_LT if on_dark else INK_2
    gradient_text(p, "Ocean Motion", grad[0], grad[1], F_DISPLAY, 22)
    run(p, " ", font=F_DISPLAY, size=22, color=grad[1])
    run(p, "ANALYTICS", font=F_MONO, size=10, color=suf_c, track=220, caps=True)
