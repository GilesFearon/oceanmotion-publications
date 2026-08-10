#!/usr/bin/env python3
"""
Ocean Motion Analytics — pitch deck starter template.

Generates a 16:9 .pptx whose colours, type and editorial treatments mirror the
website design system (styles.css, "Editorial Atlas"). One real slide per layout
type; duplicate slides in Google Slides to build a deck.

Regenerate:  python3 build_template.py
Output:      oma-pitch-template.pptx
"""

import os
from datetime import date

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ------------------------------------------------------------------ tokens
ABYSS    = RGBColor(0x05, 0x0f, 0x1c)
NAVY     = RGBColor(0x09, 0x1c, 0x33)
NAVY_MID = RGBColor(0x13, 0x29, 0x4a)
CYAN     = RGBColor(0x4f, 0xc3, 0xd7)
CYAN_LT  = RGBColor(0xa8, 0xe0, 0xea)
ACCENT   = RGBColor(0xff, 0x6a, 0x2b)
PAPER    = RGBColor(0xf4, 0xef, 0xe3)
PAPER_2  = RGBColor(0xef, 0xe9, 0xd9)
INK      = RGBColor(0x0b, 0x14, 0x20)
INK_2    = RGBColor(0x4a, 0x56, 0x65)
INK_3    = RGBColor(0x8a, 0x94, 0xa2)
RULE     = RGBColor(0xd6, 0xcd, 0xb6)

F_DISPLAY = "Instrument Serif"
F_BODY    = "IBM Plex Sans"
F_MONO    = "IBM Plex Mono"

# Eddy logo mark. PNGs are rasterised from the website glyphs (assets/*.svg,
# mirrored from oceanmotion-web) so the deck stays self-contained. To refresh:
#   convert -background none -density 1600 assets/om-mark-dark.svg \
#           -resize 1024x1024 assets/om-mark-dark.png   (and -light likewise)
ASSET_DIR  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
MARK_DARK  = os.path.join(ASSET_DIR, "om-mark-dark.png")   # cyan + accent, dark grounds
MARK_LIGHT = os.path.join(ASSET_DIR, "om-mark-light.png")  # navy + accent, light grounds

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.92)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def set_theme_fonts(presentation, major, minor):
    """Point the theme's major/minor font scheme at the brand faces, so any new
    text box added later inherits them instead of the default Calibri."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    part = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    theme = etree.fromstring(part.blob)
    scheme = theme.find(qn('a:themeElements')).find(qn('a:fontScheme'))
    for tag, face in ((qn('a:majorFont'), major), (qn('a:minorFont'), minor)):
        scheme.find(tag).find(qn('a:latin')).set('typeface', face)
    part._blob = etree.tostring(theme, xml_declaration=True,
                                encoding='UTF-8', standalone=True)


set_theme_fonts(prs, F_DISPLAY, F_BODY)


# ------------------------------------------------------------------ helpers
def slide():
    return prs.slides.add_slide(BLANK)


def flatten(r):
    """Remove the autoshape's themed <p:style> (which carries an effectRef
    shadow) so the shape renders perfectly flat, matching the website."""
    el = r._element
    style = el.find(qn('p:style'))
    if style is not None:
        el.remove(style)
    r.shadow.inherit = False
    return r


def bg(s, color):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    flatten(r)
    # push to back
    s.shapes._spTree.remove(r._element)
    s.shapes._spTree.insert(2, r._element)
    return r


def rect(s, left, top, w, h, color, line=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line; r.line.width = Pt(1)
    flatten(r)
    return r


def spacing(run, centipoints):
    """Letter-spacing via the a:rPr spc attribute (1/100 pt)."""
    run.font._rPr.set('spc', str(int(centipoints)))


def box(s, left, top, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def run(p, text, font=F_BODY, size=18, color=INK, bold=False, italic=False,
        track=None, caps=False):
    r = p.add_run()
    r.text = text.upper() if caps else text
    f = r.font
    f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color
    if track is not None:
        spacing(r, track)
    return r


def eyebrow(s, left, top, text, color=INK_2, w=Inches(6)):
    """Mono uppercase eyebrow with a short leading rule (mirrors .section-eyebrow)."""
    rule_w = Inches(0.42)
    ln = rect(s, left, top + Pt(7), rule_w, Pt(1.2), color)
    tf = box(s, left + rule_w + Inches(0.16), top, w, Inches(0.4))
    p = tf.paragraphs[0]
    run(p, text, font=F_MONO, size=12, color=color, track=180, caps=True)
    return tf


def glyph(s, left, top, height, on_dark=True):
    """Place the eddy logo mark (auto-scales width from the square source)."""
    return s.shapes.add_picture(MARK_DARK if on_dark else MARK_LIGHT,
                                left, top, height=height)


def wordmark(s, left, top, on_dark=True):
    # eddy mark, then name + suffix — mirrors the site header lockup
    g_h = Inches(0.46)
    glyph(s, left - Inches(0.05), top - Inches(0.03), g_h, on_dark)
    tf = box(s, left + g_h, top, Inches(6), Inches(0.5))
    p = tf.paragraphs[0]
    name_c = PAPER if on_dark else INK
    suf_c = CYAN_LT if on_dark else INK_2
    run(p, "Ocean Motion ", font=F_DISPLAY, size=22, color=name_c)
    run(p, "ANALYTICS", font=F_MONO, size=10, color=suf_c, track=220, caps=True)


# ================================================================== 1. TITLE
s = slide(); bg(s, ABYSS)
# faint horizon rule
rect(s, MARGIN, Inches(2.4), Inches(1.6), Pt(2), ACCENT)
wordmark(s, MARGIN, Inches(0.7), on_dark=True)

tf = box(s, MARGIN, Inches(2.7), Inches(10.5), Inches(2.8))
p = tf.paragraphs[0]; p.line_spacing = 1.02
run(p, "Ocean models, built for\n", font=F_DISPLAY, size=58, color=PAPER)
p2 = tf.add_paragraph(); p2.line_spacing = 1.02
run(p2, "decisions ", font=F_DISPLAY, size=58, color=PAPER)
run(p2, "that can't wait", font=F_DISPLAY, size=58, color=ACCENT, italic=True)
run(p2, ".", font=F_DISPLAY, size=58, color=PAPER)

tf = box(s, MARGIN, Inches(5.9), Inches(9), Inches(1))
p = tf.paragraphs[0]
run(p, "PITCH DECK", font=F_MONO, size=11, color=CYAN_LT, track=200, caps=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(6)
run(p2, f"Prepared for [Client]  ·  {date.today():%B %Y}", font=F_BODY, size=14, color=INK_3)


# ================================================================== 2. CONTENTS
s = slide(); bg(s, PAPER)
eyebrow(s, MARGIN, Inches(0.95), "Contents")
tf = box(s, MARGIN, Inches(1.4), Inches(9), Inches(1))
run(tf.paragraphs[0], "What we'll cover.", font=F_DISPLAY, size=44, color=INK)

items = ["The problem you're solving",
         "How we model it",
         "Live forecast — your site",
         "What you'd receive, day to day",
         "Scope, timeline & next steps"]
tf = box(s, MARGIN, Inches(2.9), Inches(10), Inches(4))
for i, it in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(14)
    run(p, f"{i+1:02d}", font=F_MONO, size=15, color=ACCENT, track=80)
    run(p, "    " + it, font=F_BODY, size=22, color=INK)


# ================================================================== 3. DIVIDER
s = slide(); bg(s, NAVY)
eyebrow(s, MARGIN, Inches(3.0), "Section 01", color=CYAN)
tf = box(s, MARGIN, Inches(3.5), Inches(11), Inches(2))
p = tf.paragraphs[0]; p.line_spacing = 1.03
run(p, "How we ", font=F_DISPLAY, size=60, color=PAPER)
run(p, "model", font=F_DISPLAY, size=60, color=ACCENT, italic=True)
run(p, " your water.", font=F_DISPLAY, size=60, color=PAPER)


# ================================================================== 4. CONTENT
s = slide(); bg(s, PAPER)
eyebrow(s, MARGIN, Inches(0.95), "Approach")
tf = box(s, MARGIN, Inches(1.4), Inches(10.5), Inches(1.4))
p = tf.paragraphs[0]; p.line_spacing = 1.04
run(p, "Physics first, ", font=F_DISPLAY, size=42, color=INK)
run(p, "tuned to your site", font=F_DISPLAY, size=42, color=ACCENT, italic=True)
run(p, ".", font=F_DISPLAY, size=42, color=INK)

paras = [
    "A regional ocean model nested down to the port, plant or site of interest.",
    "Forced with the best available atmospheric, tidal and boundary data, updated every day.",
    "Outputs framed as the decision you actually make — operability windows, go/no-go, thresholds.",
]
tf = box(s, MARGIN, Inches(3.1), Inches(8.4), Inches(3.4))
for i, t in enumerate(paras):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(16); p.line_spacing = 1.4
    run(p, "—  ", font=F_BODY, size=19, color=ACCENT)
    run(p, t, font=F_BODY, size=19, color=INK_2)


# ================================================================== 5. TWO COLUMN
s = slide(); bg(s, PAPER)
eyebrow(s, MARGIN, Inches(0.95), "What we do")
tf = box(s, MARGIN, Inches(1.4), Inches(10), Inches(1))
run(tf.paragraphs[0], "From physics to decisions.", font=F_DISPLAY, size=42, color=INK)

cards = [
    ("Recurring", "We run them every day",
     "Regional and hyper-local forecasts at your site, updated daily. "
     "Operability windows, go/no-go calls, custom thresholds."),
    ("One-off", "We build ocean models",
     "Model development, hindcasts, dispersion and transport, design criteria. "
     "Each engagement scoped to a specific question."),
]
cw = Inches(5.55); gap = Inches(0.5)
cx = [MARGIN, MARGIN + cw + gap]
for (tag, head, body_t), x in zip(cards, cx):
    rect(s, x, Inches(2.9), cw, Inches(3.4), PAPER_2)
    rect(s, x, Inches(2.9), cw, Pt(3), ACCENT)
    inset = Inches(0.45)
    tf = box(s, x + inset, Inches(3.25), cw - inset * 2, Inches(0.4))
    run(tf.paragraphs[0], tag, font=F_MONO, size=11, color=ACCENT, track=180, caps=True)
    tf = box(s, x + inset, Inches(3.75), cw - inset * 2, Inches(0.7))
    run(tf.paragraphs[0], head, font=F_DISPLAY, size=28, color=INK)
    tf = box(s, x + inset, Inches(4.6), cw - inset * 2, Inches(1.6))
    p = tf.paragraphs[0]; p.line_spacing = 1.4
    run(p, body_t, font=F_BODY, size=15, color=INK_2)


# ================================================================== 6. BIG STATEMENT
s = slide(); bg(s, PAPER)
tf = box(s, MARGIN, 0, Inches(11.5), EMU_H, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT; p.line_spacing = 1.06
run(p, "The people who have to know\nwhat the sea is doing ", font=F_DISPLAY, size=46, color=INK)
run(p, "shouldn't\nbe guessing", font=F_DISPLAY, size=46, color=ACCENT, italic=True)
run(p, ".", font=F_DISPLAY, size=46, color=INK)


# ================================================================== 7. DEMO / FULL-BLEED
s = slide(); bg(s, ABYSS)
eyebrow(s, MARGIN, Inches(0.7), "Live forecast", color=CYAN)
tf = box(s, MARGIN, Inches(1.12), Inches(10), Inches(0.7))
run(tf.paragraphs[0], "Your site, today.", font=F_DISPLAY, size=34, color=PAPER)
# screenshot placeholder frame
ph = rect(s, MARGIN, Inches(2.05), Inches(11.5), Inches(4.7), NAVY_MID)
ph.line.color.rgb = CYAN; ph.line.width = Pt(1)
tf = box(s, MARGIN, Inches(4.1), Inches(11.5), Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "[ Replace with dashboard screenshot / recording ]",
    font=F_MONO, size=12, color=CYAN_LT, track=120, caps=True)


# ================================================================== 8. METRICS
s = slide(); bg(s, NAVY)
eyebrow(s, MARGIN, Inches(0.95), "What you'd receive", color=CYAN)
tf = box(s, MARGIN, Inches(1.4), Inches(10), Inches(1))
run(tf.paragraphs[0], "Every single day.", font=F_DISPLAY, size=42, color=PAPER)

metrics = [("Daily", "forecast cycle"),
           ("72 h", "lead time"),
           ("Your", "thresholds")]
mw = Inches(3.6)
for i, (big, lbl) in enumerate(metrics):
    x = MARGIN + i * (mw + Inches(0.2))
    tf = box(s, x, Inches(3.3), mw, Inches(1.3))
    run(tf.paragraphs[0], big, font=F_DISPLAY, size=66, color=ACCENT)
    tf = box(s, x, Inches(4.7), mw, Inches(0.6))
    run(tf.paragraphs[0], lbl, font=F_MONO, size=12, color=CYAN_LT, track=150, caps=True)


# ================================================================== 9. CLOSING
s = slide(); bg(s, ABYSS)
wordmark(s, MARGIN, Inches(0.7), on_dark=True)
rect(s, MARGIN, Inches(2.5), Inches(1.6), Pt(2), ACCENT)
tf = box(s, MARGIN, Inches(2.8), Inches(11), Inches(1.6))
p = tf.paragraphs[0]; p.line_spacing = 1.03
run(p, "Tell us what you're\n", font=F_DISPLAY, size=56, color=PAPER)
run(p, "working on", font=F_DISPLAY, size=56, color=ACCENT, italic=True)
run(p, ".", font=F_DISPLAY, size=56, color=PAPER)

tf = box(s, MARGIN, Inches(5.4), Inches(10), Inches(1.2))
p = tf.paragraphs[0]
run(p, "Giles Fearon", font=F_BODY, size=18, color=PAPER, bold=True)
p2 = tf.add_paragraph(); p2.space_before = Pt(4)
run(p2, "giles@oceanmotionanalytics.com", font=F_MONO, size=13, color=CYAN_LT)
p3 = tf.add_paragraph(); p3.space_before = Pt(2)
run(p3, "oceanmotionanalytics.com", font=F_MONO, size=13, color=INK_3)


prs.save("oma-pitch-template.pptx")
print("wrote oma-pitch-template.pptx —", len(prs.slides._sldIdLst), "slides")
